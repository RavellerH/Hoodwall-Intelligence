#!/usr/bin/env python3
"""Import wallet/entity analysis from xlsx workbooks into knowledge-base records.

Built for the shapes actually produced by this project's analysis exports:

  Smart Money Wallets    masked EVM wallets + handle + win rate    -> wallets + entities
  Hyperliquid Clusters   full EVM addresses + perps positions      -> wallets
  Named Individuals      a person with a disclosed address         -> entities + wallets
  BTC Institutional      an entity with holdings but NO address    -> entities (holdings)

Nothing is written unless --write is passed: the default is a dry run that
prints what would change, because these files are hand-curated and an
importer silently overwriting an afternoon of analysis is unacceptable.

    python scripts/kb_import.py <file.xlsx> [...] [--write]
"""
import argparse
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

import yaml

ROOT = Path(__file__).resolve().parent.parent
KNOWLEDGE = ROOT / "knowledge"

FULL_EVM = re.compile(r"^0x[a-fA-F0-9]{40}$")
MASKED = re.compile(r"^(0x[a-fA-F0-9]{3,8})\s*(?:\.{2,3}|[…⋯])\s*([a-fA-F0-9]{3,8})$")

ENTITY_TYPE_MAP = {
    "exchange": "exchange", "etf issuer": "etf", "corporate treasury": "treasury",
    "corporate / miner": "treasury", "state holder": "treasury",
    "founder / dormant": "individual",
}


def slugify(text):
    slug = re.sub(r"[^a-z0-9]+", "-", str(text).lower()).strip("-")
    return slug[:60] or "unnamed"


def clean(value):
    return str(value).strip() if value is not None else ""


def parse_number(value):
    """Pull the first number out of '~13,000 BTC' / '+$68,468.96' / '40.0%'."""
    if value is None:
        return None
    text = str(value).replace(",", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    number = float(match.group())
    return -abs(number) if text.strip().startswith("-") or "-$" in text else number


def sheet_rows(ws):
    """Yield dict rows keyed by header, skipping blank rows."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [clean(h) for h in rows[0]]
    out = []
    for raw in rows[1:]:
        if not any(v is not None and clean(v) for v in raw):
            continue
        out.append({h: raw[i] for i, h in enumerate(headers) if h})
    return out


# --- per-sheet importers --------------------------------------------------

def import_smart_money(ws, source_key, feed_narrative=None):
    """Masked memecoin wallets. Handles become entities."""
    wallets, entities = [], {}
    for row in sheet_rows(ws):
        raw = clean(row.get("Wallet (masked)") or row.get("Wallet"))
        if not raw:
            continue
        handle = clean(row.get("Handle"))
        entity_key = slugify(handle.lstrip("@")) if handle else None

        # "OG($$242.22); SOCIAL($$241.05)" -> ["OG", "SOCIAL"]
        bought = clean(row.get("Tokens Bought (this feed)"))
        tokens = sorted({
            t.strip() for t in re.findall(r"([^;()]+)\(", bought) if t.strip()
        }) if bought else []

        record = {"chain": "robinhood", "source": source_key,
                  "labels": ["smart_money"], "confidence": "reported",
                  "tokens": tokens,
                  "narratives": [feed_narrative] if feed_narrative else []}
        if FULL_EVM.match(raw):
            record["address"] = raw.lower()
        elif MASKED.match(raw):
            record["masked"] = raw.lower()
        else:
            continue

        win = parse_number(row.get("Reported Win Rate"))
        pnl = parse_number(row.get("Reported Realized PnL"))
        note = []
        if win is not None:
            note.append(f"reported win rate {win:.1f}%")
        if pnl is not None:
            note.append(f"reported realized PnL ${pnl:,.0f}")
        appearances = parse_number(row.get("# Appearances in feed") or row.get("# Appearances"))
        if appearances:
            note.append(f"{int(appearances)} feed appearance(s)")
        if note:
            # Explicitly vendor-reported: these are claims, not observations.
            record["notes"] = "Vendor-reported: " + "; ".join(note)
        if handle:
            record["handle"] = handle
            record["entity"] = entity_key
            entities[entity_key] = {
                "key": entity_key, "name": handle.lstrip("@"), "type": "individual",
                "confidence": "reported", "handles": {"twitter": handle},
                "notes": f"Tagged as smart money by {source_key}."
                         + (f" Vendor-reported win rate {win:.1f}%." if win is not None else ""),
            }
        wallets.append(record)
    return wallets, entities


def import_hyperliquid(ws):
    """Live leaderboard wallets with real positions."""
    wallets = []
    for row in sheet_rows(ws):
        address = clean(row.get("Wallet"))
        if not FULL_EVM.match(address):
            continue
        value = parse_number(row.get("Account Value ($)"))
        pnl = parse_number(row.get("Unrealized PnL ($)"))
        cluster = clean(row.get("Asset Cluster"))
        direction = clean(row.get("Direction"))
        note = []
        if value:
            note.append(f"account value ${value:,.0f}")
        if pnl is not None:
            note.append(f"unrealized PnL ${pnl:,.0f}")
        if cluster and cluster != "--":
            note.append(f"concentrated in {cluster}")
        if direction and direction != "--":
            note.append(f"net {direction.lower()}")
        wallets.append({
            "chain": "hyperliquid", "address": address.lower(),
            "labels": ["whale", "perps_trader"], "confidence": "confirmed",
            "notes": "Leaderboard snapshot: " + "; ".join(note) if note else None,
        })
    return [{k: v for k, v in w.items() if v is not None} for w in wallets]


def import_named_individuals(ws):
    """A person with a disclosed address: entity + its wallets."""
    wallets, entities = [], {}
    for row in sheet_rows(ws):
        name = clean(row.get("Name/Handle"))
        if not name:
            continue
        key = slugify(re.split(r"[ (]", name)[0])
        raw = clean(row.get("Wallet") or row.get("Known wallet / identifier"))
        address = raw.split()[0] if raw else ""
        note = clean(row.get("Note") or row.get("Profile"))
        chains_seen = clean(row.get("Chain(s)") or row.get("Chain(s) seen"))

        entities[key] = {
            "key": key, "name": name, "type": "individual",
            # Self-disclosed and publicly attributed, unlike the vendor tags.
            "confidence": "confirmed",
            "handles": {"twitter": m.group()} if (m := re.search(r"@\w+", name)) else {},
            "notes": note[:400] if note else None,
            "chains_seen": chains_seen or None,
        }
        entities[key] = {k: v for k, v in entities[key].items() if v not in (None, {}, "")}

        if FULL_EVM.match(address):
            wallets.append({
                "chain": "hyperliquid" if "hyperliquid" in chains_seen.lower() else "ethereum",
                "address": address.lower(), "entity": key,
                "labels": ["smart_money"], "confidence": "confirmed",
                "notes": f"Publicly attributed to {name}.",
            })
    return wallets, entities


def import_btc_entities(ws):
    """Entities tracked by holdings, not addresses.

    A BTC treasury's coins are spread across many addresses that are not
    individually interesting, so the entity carries the holding directly and
    has no wallet records at all.
    """
    entities = {}
    for row in sheet_rows(ws):
        name = clean(row.get("Entity"))
        if not name or name.lower().startswith("source:"):
            continue
        held = clean(row.get("Approx. BTC held"))
        kind = clean(row.get("Type")).lower()
        key = slugify(re.split(r"[(]", name)[0])
        amount = parse_number(held)
        entities[key] = {
            "key": key, "name": re.split(r"\s*\(", name)[0].strip(),
            "type": ENTITY_TYPE_MAP.get(kind, "unknown"),
            "confidence": "reported",
            "holdings": [{"chain": "bitcoin", "asset": "BTC",
                          "amount": amount, "as_reported": held}],
            "notes": f"{kind.title()}. Holdings are entity-level; coins span many addresses.",
        }
    return entities


SHEET_HANDLERS = {
    "Smart Money Wallets": "smart_money",
    "Memecoin Smart Money": "smart_money",
    "Hyperliquid Clusters": "hyperliquid",
    "Hyperliquid Top Traders": "hyperliquid",
    "Named Individuals": "named",
    "Notable Named Wallets": "named",
    "BTC Institutional": "btc",
    "BTC Whale Entities": "btc",
}


def merge_yaml_list(path, new_rows, dedupe_key, write):
    """Merge rows into a YAML list file without clobbering existing entries."""
    existing = []
    if path.exists():
        loaded = yaml.safe_load(path.read_text(encoding="utf-8")) or []
        existing = loaded if isinstance(loaded, list) else []
    seen = {r.get(dedupe_key) or r.get("masked") for r in existing}
    added = [r for r in new_rows if (r.get(dedupe_key) or r.get("masked")) not in seen]
    if not added:
        return 0
    if write:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(
            yaml.safe_dump(existing + added, sort_keys=False, allow_unicode=True, width=100),
            encoding="utf-8")
    return len(added)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("files", nargs="+", help="xlsx workbooks to import")
    ap.add_argument("--write", action="store_true", help="actually write (default: dry run)")
    ap.add_argument("--source", default="intel-hood-vantis", help="source key for feed wallets")
    ap.add_argument("--narrative", default=None,
                    help="narrative key to link imported feed wallets to")
    args = ap.parse_args()

    all_wallets, all_entities = {}, {}

    for filename in args.files:
        path = Path(filename)
        if not path.exists():
            print(f"! {path}: not found")
            continue
        wb = openpyxl.load_workbook(path, data_only=True)
        print(f"\n{path.name}")
        for sheet_name in wb.sheetnames:
            kind = SHEET_HANDLERS.get(sheet_name)
            if not kind:
                print(f"  - {sheet_name}: no handler, skipped")
                continue
            ws = wb[sheet_name]
            if kind == "smart_money":
                wallets, entities = import_smart_money(
                    ws, args.source, args.narrative)
            elif kind == "hyperliquid":
                wallets, entities = import_hyperliquid(ws), {}
            elif kind == "named":
                wallets, entities = import_named_individuals(ws)
            else:
                wallets, entities = [], import_btc_entities(ws)

            for w in wallets:
                all_wallets.setdefault(f"{w['chain']}:{w.get('address') or w.get('masked')}", w)
            all_entities.update(entities)
            print(f"  + {sheet_name}: {len(wallets)} wallet(s), {len(entities)} entity(ies)")

    by_chain = {}
    for w in all_wallets.values():
        by_chain.setdefault(w["chain"], []).append(w)

    print("\n" + "=" * 60)
    print("DRY RUN - nothing written (pass --write to apply)" if not args.write
          else "WRITING")
    total = 0
    for chain, rows in sorted(by_chain.items()):
        n = merge_yaml_list(KNOWLEDGE / "wallets" / f"{chain}.yml", rows, "address", args.write)
        total += n
        print(f"  wallets/{chain}.yml: +{n} new (of {len(rows)} seen)")
    for key, entity in sorted(all_entities.items()):
        n = merge_yaml_list(KNOWLEDGE / "entities" / f"{key}.yml", [entity], "key", args.write)
        total += n
    print(f"  entities/: +{sum(1 for _ in all_entities)} seen")
    print(f"\n{total} record(s) {'written' if args.write else 'would be written'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
